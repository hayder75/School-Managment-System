#!/usr/bin/env python3
"""Normalize the school's real data files into JSON for seeding.
Reads from `new real datat/` and writes `new real datat/seed/` (gitignored)."""
import os, re, json, collections, zipfile
import openpyxl

BASE = "/home/hayder/sms/new real datat"
OUT = os.path.join(BASE, "seed")
os.makedirs(OUT, exist_ok=True)

def norm(s):
    return re.sub(r"[^a-z0-9 ]", "", re.sub(r"\s+", " ", (s or "").strip().lower()))

def phone_norm(p):
    d = re.sub(r"\D", "", str(p or ""))
    if d.startswith("251") and len(d) == 12:
        d = "0" + d[3:]
    if len(d) == 9:
        d = "0" + d
    return d

def class_meta(code):
    c = (code or "").strip()
    u = c.upper().replace("NURSURY", "NURSERY").replace("NURSERY", "NURSERY")
    m = re.match(r"^NURSERY\s*([A-F])$", u)
    if m:
        return {"level_group": "nursery", "grade_level": 0, "section": m.group(1), "name": f"Nursery {m.group(1)}"}
    m = re.match(r"^LKG\s*([A-F])$", u)
    if m:
        return {"level_group": "kg", "grade_level": 1, "section": m.group(1), "name": f"LKG {m.group(1)}"}
    m = re.match(r"^UKG\s*([A-F])$", u)
    if m:
        return {"level_group": "kg", "grade_level": 2, "section": m.group(1), "name": f"UKG {m.group(1)}"}
    m = re.match(r"^GRADE\s*(\d+)\s*([A-F])$", u)
    if m:
        g = int(m.group(1)); s = m.group(2)
        return {"level_group": "primary", "grade_level": g, "section": s, "name": f"Grade {g} {s}"}
    return None

# ---- Master students ----
def g(r, i):
    return (str(r[i]).strip() if len(r) > i and r[i] is not None else "")

wb = openpyxl.load_workbook(os.path.join(BASE, "All Student name list Sep 14.xlsx"), data_only=True, read_only=True)
rows = list(wb["Sheet1"].iter_rows(values_only=True))[1:]
wb.close()

students = []
seen_names = set()
for r in rows:
    if not (r and (r[2] or r[1])): continue
    reg = g(r, 1)
    cm = class_meta(g(r, 5))
    archived = g(r, 5) == "__ARCHIVE__"
    st = {
        "reg": reg if re.match(r"^M", reg, re.I) else "",
        "first": g(r, 2), "middle": g(r, 3), "last": g(r, 4),
        "class_code": g(r, 5), "class": cm, "phone": phone_norm(g(r, 6)),
        "gender": {"M": "male", "F": "female"}.get(g(r, 7), ""),
        "archived": archived,
        "source": "master",
    }
    students.append(st)
    seen_names.add(norm(f"{st['first']} {st['middle']} {st['last']}"))

# ---- KG extra students (not in master) ----
extras = []
for f, idx in [("Nursery_Student_List_2019.xlsx", None), ("LKG -Student List 2019.xlsx", None), ("UKG_Student List_2019_.xlsx", None)]:
    wb = openpyxl.load_workbook(os.path.join(BASE, f), data_only=True, read_only=True)
    for ws in wb.worksheets:
        if ws.title.startswith("All"): continue
        # section -> class meta
        code = ws.title
        if f.startswith("Nursery"): code = "NURSERY " + ws.title.split()[-1]
        cm = class_meta(code)
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if not r or len(r) < 4: continue
            first, mid, last = g(r, 1), g(r, 2), g(r, 3)
            if not first or first.lower() in ("first name", "s.no"): continue
            nm = norm(f"{first} {mid} {last}")
            if nm in seen_names: continue
            seen_names.add(nm)
            def col(j):
                return g(r, j) if len(r) > j else ""
            extras.append({
                "first": first, "middle": mid, "last": last,
                "class": cm, "gender": {"M": "male", "F": "female"}.get(col(4), ""),
                "age": col(5), "mother_phone": phone_norm(col(6)), "father_phone": phone_norm(col(7)),
                "remark": col(8), "source": f,
            })
    wb.close()

# ---- Staff ----
wb = openpyxl.load_workbook(os.path.join(BASE, "All Staff Sep 2026.xlsx"), data_only=True, read_only=True)
staff = []
for i, r in enumerate(wb["Sheet1"].iter_rows(values_only=True)):
    if i < 3 or not r or not r[1]: continue
    if str(r[1]).strip() == "Full Name": continue
    staff.append({
        "name": str(r[1]).strip(), "job_title": g(r, 2), "joined": g(r, 3),
        "salary": g(r, 4), "allowance": g(r, 5),
        "gender": {"male": "male", "female": "female"}.get(str(r[6]).strip().lower(), g(r, 6).lower()),
        "role": g(r, 7), "phone": phone_norm(g(r, 8)), "dob": g(r, 9),
        "qualification": g(r, 10), "field": g(r, 11),
    })
wb.close()

# ---- Subjects ----
wb = openpyxl.load_workbook(os.path.join(BASE, "Subjects and Levels.xlsx"), data_only=True, read_only=True)
subjects = []
for i, r in enumerate(wb["Sheet1"].iter_rows(values_only=True)):
    if i < 3 or not r or not r[1]: continue
    subjects.append({"name": str(r[1]).strip(), "range": str(r[2] or "").strip(), "note": str(r[3] or "").strip()})
wb.close()

# ---- Homeroom ----
HOMEROOM = """Nigist Tema - 1B
Betrework Belete - 4B
Berehanu Naremo - 2C
Yalew Solomon - 3C
Dawit Ermiyas - 5E
Mekdes Ayele - 6A
Ayele Alemayehu - 1E
Henok Dagne - 2D
Akalu Girma - 6E
Joseph Muluneh - 2F
Bereket Dilnesaw - 3B
Aster Tsegaye - 4E
Awol Muzemil - 6B
Dereje Degualem - 4A
Mengistu Alemayehu - 2E
Mitike Yetera - 3A
Bernabas Dibaba - 4D
Biruktesfa Getachew - 7C
Wondesen Dejene - 1D
Yehualashet Gadisa - 3F
Samuel Bogale - 1A
Million Birhanu - 6C
Kassahun Kanchula - 3G
Aklilu Chufamo - 5B
Andualem Belay - 5D
Tigabu Delelegn - 8B
Workinesh Berhanu - 2G
Mathewos Elias - 5C
Fitsum Tadele - 8D
Mesay F/Yesus - 4C
Unassigned - 3E
Seleshi Kura - 5A
Philipo Tamru - 6D
Berhanu Kebede - 8E
Asefa Abebo - 7B
Amanuel Sodano - 8A
Dereje Mekonnen - 7E
Jonny Hailu - 1C
Yohannes Assefa - 7D
Eyerusalem Argaw - 7A
Biniyam Wondimu - 8C
Abayneh Nushe - 2A
Unassigned - 2B
Unassigned - 1F
Unassigned - 3D"""
homeroom = []
for line in HOMEROOM.strip().split("\n"):
    t, c = line.split(" - ", 1)
    homeroom.append({"teacher": t.strip(), "class_code": c.strip()})

def dump(name, obj):
    with open(os.path.join(OUT, name), "w") as fh:
        json.dump(obj, fh, ensure_ascii=False, indent=1)
    print(f"{name}: {len(obj) if isinstance(obj, list) else 'obj'}")

dump("students_master.json", students)
dump("students_extra.json", extras)
dump("staff.json", staff)
dump("subjects.json", subjects)
dump("homeroom.json", homeroom)
print("done")
